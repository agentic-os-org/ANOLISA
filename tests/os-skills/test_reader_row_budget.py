"""Bound actual table IO and label every report's analyzed-row scope."""

import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import pandas as pd
from openpyxl import Workbook

SCRIPT = Path(__file__).parents[2] / "src/os-skills/others/xlsx/scripts/xlsx_reader.py"


class ReaderRowBudgetTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        spec = importlib.util.spec_from_file_location("budget_table_reader", SCRIPT)
        cls.reader = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.reader)

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.xlsx = self.root / "data.xlsx"
        workbook = Workbook()
        for index, (title, count) in enumerate((("Long", 6), ("短表", 2), ("Empty", 0))):
            sheet = workbook.active if index == 0 else workbook.create_sheet()
            sheet.title = title
            sheet.append(["label", "amount"])
            for row in range(count):
                sheet.append([f'收入, "{row}"', row + 1])
        workbook.save(self.xlsx)
        workbook.close()
        self.csv = self.root / "export.csv"
        pd.DataFrame(
            {
                "label": ['one, "quoted"', "two\nlines", "三", "four", "five", "six"],
                "amount": [1, 2, 3, 4, 5, 6],
            }
        ).to_csv(self.csv, index=False)
        self.tsv = self.root / "export.tsv"
        pd.read_csv(self.csv).to_csv(self.tsv, index=False, sep="\t")

    def run_cli(self, path, *arguments):
        return subprocess.run(
            [sys.executable, str(SCRIPT), str(path), *arguments],
            env={**os.environ, "PYTHONUTF8": "1"},
            capture_output=True,
            text=True,
            encoding="utf-8",
            timeout=30,
        )

    def test_actual_all_sheets_limit_independently_with_truncation_metadata(self):
        sheets = self.reader.detect_and_load(str(self.xlsx), max_rows=3)
        self.assertEqual(list(sheets), ["Long", "短表", "Empty"])
        self.assertEqual([len(frame) for frame in sheets.values()], [3, 2, 0])
        self.assertTrue(sheets["Long"].attrs["reader_truncated"])
        self.assertFalse(sheets["短表"].attrs["reader_truncated"])
        self.assertFalse(sheets["Empty"].attrs["reader_truncated"])
        self.assertEqual(sheets["Long"]["amount"].tolist(), [1, 2, 3])

    def test_named_sheet_retains_selection_and_limit(self):
        sheets = self.reader.detect_and_load(str(self.xlsx), sheet_name_filter="Long", max_rows=2)
        self.assertEqual(list(sheets), ["Long"])
        self.assertEqual(sheets["Long"]["amount"].tolist(), [1, 2])

    def test_flat_formats_count_data_records_not_physical_lines(self):
        for path in (self.csv, self.tsv):
            with self.subTest(path=path):
                frame = self.reader.detect_and_load(str(path), max_rows=2)["export"]
                self.assertEqual(frame["amount"].tolist(), [1, 2])
                self.assertEqual(frame.iloc[1]["label"], "two\nlines")
                self.assertTrue(frame.attrs["reader_truncated"])
                self.assertEqual(frame._reader_encoding, "utf-8-sig")

    def test_lookahead_row_does_not_change_analyzed_prefix_types(self):
        self.csv.write_text("amount\n1\n2\n3\noutside-text\n", encoding="utf-8")
        self.tsv.write_text("amount\n1\n2\n3\noutside-text\n", encoding="utf-8")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Long"
        for value in ("amount", 1, 2, 3, "outside-text"):
            sheet.append([value])
        workbook.save(self.xlsx)
        workbook.close()
        for path in (self.csv, self.tsv, self.xlsx):
            with self.subTest(path=path):
                sheets = self.reader.detect_and_load(str(path), max_rows=3)
                frame = next(iter(sheets.values()))
                self.assertEqual(frame["amount"].tolist(), [1, 2, 3])
                self.assertTrue(pd.api.types.is_numeric_dtype(frame["amount"]))
                self.assertEqual(
                    next(iter(self.reader.compute_stats(sheets).values()))["amount"]["mean"], 2
                )
                self.assertTrue(frame.attrs["reader_truncated"])

    def test_excel_budget_is_pushed_into_pandas_io(self):
        with mock.patch.object(pd, "read_excel", wraps=pd.read_excel) as read:
            self.reader.detect_and_load(str(self.xlsx), max_rows=3)
        self.assertEqual(read.call_args_list[0].kwargs["nrows"], 4)

    def test_flat_budget_is_pushed_into_pandas_io(self):
        with mock.patch.object(pd, "read_csv", wraps=pd.read_csv) as read:
            self.reader.detect_and_load(str(self.csv), max_rows=3)
        self.assertEqual(read.call_args_list[0].kwargs["nrows"], 4)

    def test_exact_limit_is_not_marked_truncated(self):
        sheets = self.reader.detect_and_load(str(self.xlsx), max_rows=2)
        self.assertFalse(sheets["短表"].attrs["reader_truncated"])
        self.assertEqual(sheets["短表"].attrs["reader_row_limit"], 2)

    def test_invalid_api_limits_are_rejected(self):
        for value in (0, -1, 1.5, "2", True):
            with self.subTest(value=value), self.assertRaises(ValueError):
                self.reader.detect_and_load(str(self.xlsx), max_rows=value)

    def test_actual_cli_json_reports_scope_and_loaded_rows(self):
        result = self.run_cli(self.xlsx, "--max-rows", "3", "--json")
        self.assertEqual(result.returncode, 0, result.stderr)
        report = json.loads(result.stdout)
        self.assertEqual(report["sampling"]["scope"], "analyzed_rows_only")
        self.assertEqual(report["sampling"]["max_rows_per_sheet"], 3)
        self.assertEqual(
            report["sampling"]["sheets"]["Long"], {"rows_analyzed": 3, "truncated": True}
        )
        self.assertEqual(report["structure"]["Long"]["shape"]["rows"], 3)
        self.assertEqual(report["stats"]["Long"]["amount"]["count"], 3)
        self.assertEqual(report["stats"]["Long"]["amount"]["mean"], 2)

    def test_actual_cli_quality_only_remains_compatible_and_scoped(self):
        result = self.run_cli(
            self.xlsx, "--max-rows", "2", "--quality", "--json", "--sheet", "Long"
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        report = json.loads(result.stdout)
        self.assertEqual(report["stats"], {})
        self.assertEqual(list(report["sampling"]["sheets"]), ["Long"])
        self.assertTrue(report["sampling"]["sheets"]["Long"]["truncated"])

    def test_actual_human_report_never_presents_sample_as_file_totals(self):
        result = self.run_cli(self.csv, "--max-rows", "2")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Analyzed rows", result.stdout)
        self.assertIn("first 2", result.stdout)
        self.assertIn("analyzed rows only", result.stdout)
        self.assertIn("additional rows", result.stdout.lower())
        self.assertNotIn("Total rows across all sheets", result.stdout)

    def test_default_json_schema_and_full_data_remain_unchanged(self):
        sheets = self.reader.detect_and_load(str(self.xlsx))
        self.assertEqual(len(sheets["Long"]), 6)
        self.assertNotIn("reader_row_limit", sheets["Long"].attrs)
        result = self.run_cli(self.csv, "--json")
        self.assertEqual(result.returncode, 0, result.stderr)
        report = json.loads(result.stdout)
        self.assertEqual(set(report), {"file", "structure", "quality", "stats"})
        self.assertEqual(report["structure"]["export"]["shape"]["rows"], 6)

    def test_cli_invalid_limits_are_usage_errors(self):
        for value in ("0", "-1", "x", "1.5"):
            with self.subTest(value=value):
                result = self.run_cli(self.csv, "--max-rows", value)
                self.assertEqual(result.returncode, 2, result.stderr)
                self.assertIn("positive integer", result.stderr)
                self.assertNotIn("Traceback", result.stderr)

    def test_bounded_analysis_preserves_all_original_bytes(self):
        for path in (self.xlsx, self.csv, self.tsv):
            with self.subTest(path=path):
                original = path.read_bytes()
                self.reader.detect_and_load(str(path), max_rows=2)
                self.assertEqual(path.read_bytes(), original)

    def test_header_only_table_has_zero_rows_and_no_truncation(self):
        self.csv.write_text("label,amount\n", encoding="utf-8")
        frame = self.reader.detect_and_load(str(self.csv), max_rows=2)["export"]
        self.assertEqual(list(frame.columns), ["label", "amount"])
        self.assertTrue(frame.empty)
        self.assertFalse(frame.attrs["reader_truncated"])


if __name__ == "__main__":
    unittest.main()
