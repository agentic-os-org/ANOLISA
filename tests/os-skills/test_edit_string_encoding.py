"""Encode shared strings created by the insert-row and add-column editors."""

import importlib.util
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils.escape import unescape

SCRIPTS = Path(__file__).resolve().parents[2] / "src/os-skills/others/xlsx/scripts"
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


class EditorStringEncodingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.editors = []
        previous = sys.path[:]
        try:
            sys.path.insert(0, str(SCRIPTS))
            for name in ("xlsx_insert_row", "xlsx_add_column"):
                spec = importlib.util.spec_from_file_location(
                    name + "_encoding", SCRIPTS / (name + ".py")
                )
                editor = importlib.util.module_from_spec(spec)
                spec.loader.exec_module(editor)
                cls.editors.append(editor)
        finally:
            sys.path[:] = previous

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        (self.root / "xl").mkdir()
        self.path = self.root / "xl/sharedStrings.xml"

    def initialize(self, encoded="existing"):
        self.path.write_text(
            f'<sst xmlns="{NS}" count="1" uniqueCount="1"><si><t>{encoded}</t></si></sst>',
            encoding="utf-8",
        )

    def texts(self):
        return [x.text or "" for x in ET.parse(self.path).getroot().iter(f"{{{NS}}}t")]

    def test_xml_illegal_characters_are_escaped_in_both_editors(self):
        values = ["before\x00after", "before\x01after", "\ufffe", "\uffff", "\ud800"]
        for editor in self.editors:
            for value in values:
                with self.subTest(editor=editor.__name__, value=repr(value)):
                    self.initialize()
                    self.assertEqual(editor.add_shared_string(str(self.root), value), 1)
                    self.assertEqual(unescape(self.texts()[1]), value)
                    self.assertEqual(ET.parse(self.path).getroot().get("uniqueCount"), "2")

    def test_literal_escape_patterns_remain_literal(self):
        values = ["Report_x000A_", "_x0041_", "_x005F_x000A_", "_x0041__x0042_", "_x00ff_"]
        for editor in self.editors:
            for value in values:
                with self.subTest(editor=editor.__name__, value=value):
                    self.initialize()
                    editor.add_shared_string(str(self.root), value)
                    encoded = self.texts()[1]
                    self.assertIn("_x005F_", encoded)
                    self.assertEqual(unescape(encoded), value)

    def test_existing_encoded_strings_reuse_semantic_indices(self):
        for editor in self.editors:
            for encoded, value in (
                ("_x0001_", "\x01"),
                ("_x005F_x0041_", "_x0041_"),
                ("_xD83D__xDE03_", "\U0001f603"),
                ("first&#13;last", "first\rlast"),
            ):
                with self.subTest(editor=editor.__name__, encoded=encoded):
                    self.initialize(encoded)
                    original = self.path.read_bytes()
                    self.assertEqual(editor.add_shared_string(str(self.root), value), 0)
                    self.assertEqual(self.path.read_bytes(), original)

    def test_existing_escape_value_does_not_alias_literal_text(self):
        for editor in self.editors:
            with self.subTest(editor=editor.__name__):
                self.initialize("_x0041_")
                self.assertEqual(editor.add_shared_string(str(self.root), "_x0041_"), 1)
                self.assertEqual(self.texts(), ["_x0041_", "_x005F_x0041_"])

    def test_carriage_returns_and_blank_lines_survive_appends(self):
        for editor in self.editors:
            with self.subTest(editor=editor.__name__):
                self.initialize("first&#13;last")
                value = "start\r\n\n   \nend"
                editor.add_shared_string(str(self.root), value)
                self.assertEqual(self.texts(), ["first\rlast", value])
                self.assertIn(b"&#13;", self.path.read_bytes())

    def test_real_packed_workbook_reads_literal_unicode_and_carriage_returns(self):
        values = ["Report_x000A_", "first\rlast", " leading & <text> ", "收入 😃"]
        for editor in self.editors:
            with self.subTest(editor=editor.__name__):
                workbook = Workbook()
                source = self.root / "source.xlsx"
                workbook.save(source)
                workbook.close()
                with zipfile.ZipFile(source) as archive:
                    archive.extractall(self.root)
                self.initialize()
                indices = [editor.add_shared_string(str(self.root), value) for value in values]
                sheet_path = self.root / "xl/worksheets/sheet1.xml"
                sheet = ET.parse(sheet_path)
                data = sheet.getroot().find(f"{{{NS}}}sheetData")
                for row, index in enumerate(indices, 1):
                    record = ET.SubElement(data, f"{{{NS}}}row", r=str(row))
                    cell = ET.SubElement(record, f"{{{NS}}}c", r=f"A{row}", t="s")
                    ET.SubElement(cell, f"{{{NS}}}v").text = str(index)
                sheet.write(sheet_path, encoding="utf-8")
                types_path = self.root / "[Content_Types].xml"
                types = ET.parse(types_path)
                ET.SubElement(
                    types.getroot(),
                    "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                    PartName="/xl/sharedStrings.xml",
                    ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
                )
                types.write(types_path, encoding="utf-8")
                relations_path = self.root / "xl/_rels/workbook.xml.rels"
                relations = ET.parse(relations_path)
                ET.SubElement(
                    relations.getroot(),
                    "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                    Id="rIdSharedStrings",
                    Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings",
                    Target="/xl/sharedStrings.xml",
                )
                relations.write(relations_path, encoding="utf-8")
                output = self.root / "output.xlsx"
                with zipfile.ZipFile(output, "w") as archive:
                    for path in self.root.rglob("*"):
                        if path.is_file() and path.suffix != ".xlsx":
                            archive.write(path, path.relative_to(self.root).as_posix())
                loaded = load_workbook(output)
                try:
                    self.assertEqual(
                        [loaded.active[f"A{row}"].value for row in range(1, len(values) + 1)],
                        values,
                    )
                finally:
                    loaded.close()

    def test_ordinary_strings_keep_reuse_and_metadata(self):
        for editor in self.editors:
            with self.subTest(editor=editor.__name__):
                self.initialize()
                original = self.path.read_bytes()
                self.assertEqual(editor.add_shared_string(str(self.root), "existing"), 0)
                self.assertEqual(self.path.read_bytes(), original)
                self.assertEqual(editor.add_shared_string(str(self.root), "new"), 1)
                self.assertEqual(self.texts(), ["existing", "new"])


if __name__ == "__main__":
    unittest.main()
