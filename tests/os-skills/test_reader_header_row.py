"""Analyze exported Excel tables whose headers follow title or blank rows."""

import contextlib
import importlib.util
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]
SCRIPT = ROOT / "src/os-skills/others/xlsx/scripts/xlsx_reader.py"


class HeaderRowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        spec = importlib.util.spec_from_file_location("reader_header_row", SCRIPT)
        cls.reader = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.reader)

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.path = self.root / "export.xlsx"
        book = Workbook()
        first = book.active
        first.title = "Sales"
        for row in (("Annual sales export",), (), ("Product", "Amount"), ("Alpha", 7), ("Beta", 3)):
            first.append(row)
        second = book.create_sheet("Returns")
        for row in (("Annual returns export",), (), ("Product", "Amount"), ("Gamma", 2)):
            second.append(row)
        book.save(self.path)
        book.close()

    def read(self, path=None, *options):
        output, errors = io.StringIO(), io.StringIO()
        with (
            mock.patch.object(
                sys, "argv", [str(SCRIPT), str(path or self.path), "--json", *options]
            ),
            contextlib.redirect_stdout(output),
            contextlib.redirect_stderr(errors),
        ):
            try:
                self.reader.main()
            except SystemExit as error:
                return error.code, output.getvalue(), errors.getvalue()
        return 0, output.getvalue(), errors.getvalue()

    def test_third_row_headers_drive_columns_preview_and_statistics(self):
        original = self.path.read_bytes()
        code, output, errors = self.read(None, "--header-row", "3")
        self.assertEqual(code, 0, errors)
        report = json.loads(output)
        self.assertEqual(report["structure"]["Sales"]["columns"], ["Product", "Amount"])
        self.assertEqual(report["structure"]["Sales"]["shape"], {"rows": 2, "cols": 2})
        self.assertEqual(
            report["structure"]["Sales"]["preview"],
            [
                {"Product": "Alpha", "Amount": 7},
                {"Product": "Beta", "Amount": 3},
            ],
        )
        self.assertEqual(report["stats"]["Sales"]["Amount"]["mean"], 5)
        self.assertEqual(self.path.read_bytes(), original)

    def test_selected_sheet_uses_its_own_header_row(self):
        code, output, errors = self.read(None, "--header-row", "3", "--sheet", "Returns")
        self.assertEqual(code, 0, errors)
        structure = json.loads(output)["structure"]
        self.assertEqual(list(structure), ["Returns"])
        self.assertEqual(structure["Returns"]["preview"], [{"Product": "Gamma", "Amount": 2}])

    def test_all_sheets_use_selected_excel_row(self):
        code, output, errors = self.read(None, "--header-row", "3")
        self.assertEqual(code, 0, errors)
        structure = json.loads(output)["structure"]
        self.assertEqual(list(structure), ["Sales", "Returns"])
        self.assertEqual(structure["Returns"]["columns"], ["Product", "Amount"])

    def test_macro_enabled_extension_can_be_read_without_changing_source(self):
        path = self.root / "export.xlsm"
        path.write_bytes(self.path.read_bytes())
        original = path.read_bytes()
        code, output, errors = self.read(path, "--header-row", "3")
        self.assertEqual(code, 0, errors)
        self.assertEqual(json.loads(output)["structure"]["Sales"]["columns"], ["Product", "Amount"])
        self.assertEqual(path.read_bytes(), original)

    def test_omitted_option_preserves_first_row_headers(self):
        path = self.root / "ordinary.xlsx"
        book = Workbook()
        book.active.append(["Product", "Amount"])
        book.active.append(["Alpha", 7])
        book.save(path)
        book.close()
        code, output, errors = self.read(path)
        self.assertEqual(code, 0, errors)
        self.assertEqual(
            json.loads(output)["structure"]["Sheet"]["preview"], [{"Product": "Alpha", "Amount": 7}]
        )

    def test_explicit_first_row_keeps_default_structure(self):
        code, output, errors = self.read(None, "--header-row", "1")
        self.assertEqual(code, 0, errors)
        default_code, default_output, default_errors = self.read()
        self.assertEqual(default_code, 0, default_errors)
        self.assertEqual(json.loads(output), json.loads(default_output))

    def test_invalid_rows_are_rejected_before_loading(self):
        for value in ("0", "-1", "abc", "1.5"):
            with self.subTest(value=value), mock.patch.object(
                self.reader, "detect_and_load"
            ) as load:
                code, output, errors = self.read(None, "--header-row", value)
                self.assertEqual(code, 2)
                self.assertIn("--header-row", errors)
                self.assertIn("positive", errors)
                self.assertEqual(output, "")
                load.assert_not_called()

    def test_csv_rejects_excel_header_option_before_loading(self):
        path = self.root / "export.csv"
        path.write_text("Product,Amount\nAlpha,7\n", encoding="utf-8")
        with mock.patch.object(self.reader, "detect_and_load") as load:
            code, output, errors = self.read(path, "--header-row", "3")
        self.assertEqual(code, 2)
        self.assertIn("Excel", errors)
        self.assertEqual(output, "")
        load.assert_not_called()


if __name__ == "__main__":
    unittest.main()
