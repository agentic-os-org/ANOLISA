#!/usr/bin/env python3
# SPDX-License-Identifier: MIT
"""Tests for style_audit.py CT_PatternFill default handling (#4990).

``patternType`` is optional in CT_PatternFill and defaults to ``none``.
openpyxl serialises that default as a bare ``<patternFill/>``, which the
audit used to read as ``''`` and compare against the literal ``'none'``,
so every openpyxl-written workbook failed with::

    [FAIL] fills[0] patternType='', must be 'none'

An absent or empty ``patternType`` must be treated as the schema default
``none``; explicit non-``none`` values keep being reported as before.
"""

import json
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))
SCRIPT = os.path.join(SCRIPTS_DIR, "style_audit.py")

try:
    import openpyxl  # noqa: F401
    HAS_OPENPYXL = True
except ImportError:  # pragma: no cover - optional end-to-end fixture
    HAS_OPENPYXL = False

# OOXML default fills, as openpyxl writes them: fills[0] omits patternType
# because "none" is the schema default.
DEFAULT_FILLS = """    <fill><patternFill/></fill>
    <fill><patternFill patternType="gray125"/></fill>"""

# Same thing with the default spelled out as an empty attribute.
EMPTY_ATTR_FILLS = """    <fill><patternFill patternType=""/></fill>
    <fill><patternFill patternType="gray125"/></fill>"""

# fills[0] carries an explicit, non-default pattern: a real violation.
SOLID_FILLS = """    <fill><patternFill patternType="solid"><fgColor rgb="FF0000FF"/><bgColor indexed="64"/></patternFill></fill>
    <fill><patternFill patternType="gray125"/></fill>"""

# fills[1] also omits patternType, i.e. defaults to "none" and so is NOT
# the required gray125 — the fix must not silence this one.
NO_GRAY125_FILLS = """    <fill><patternFill/></fill>
    <fill><patternFill/></fill>"""

# No patternFill element at all: CT_Fill allows gradientFill, and there is
# no patternType default to fall back on.
GRADIENT_FILLS = """    <fill><gradientFill/></fill>
    <fill><patternFill patternType="gray125"/></fill>"""

STYLES_TEMPLATE = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<styleSheet xmlns="{NS_MAIN}">
  <fonts count="2">
    <font><sz val="11"/><color theme="1"/></font>
    <font><sz val="11"/><color rgb="000000FF"/></font>
  </fonts>
  <fills count="2">
{{fills}}
  </fills>
  <borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>
  <cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>
  <cellXfs count="2">
    <xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>
    <xf numFmtId="0" fontId="1" fillId="0" borderId="0" xfId="0" applyFont="1"/>
  </cellXfs>
</styleSheet>
"""

WORKBOOK_XML = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="{NS_MAIN}" xmlns:r="{NS_REL}">
  <sheets>
    <sheet name="Sheet1" sheetId="1" r:id="rId1"/>
  </sheets>
</workbook>
"""

RELS_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1"
    Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"
    Target="worksheets/sheet1.xml"/>
</Relationships>
"""

# A label plus a blue-font numeric input: no colour-role, year-format or
# percentage violation, so the fills under test are the only thing that can
# move the verdict.
WORKSHEET_XML = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="{NS_MAIN}">
  <sheetData>
    <row r="1">
      <c r="A1" t="inlineStr"><is><t>Revenue</t></is></c>
      <c r="B1" s="1"><v>1234.5</v></c>
    </row>
  </sheetData>
</worksheet>
"""


def build_xlsx(path: str, fills: str) -> str:
    """Write a minimal, otherwise fully compliant workbook with the given fills."""
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml",
                   '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                   '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                   '<Default Extension="xml" ContentType="application/xml"/>'
                   '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
                   '</Types>')
        z.writestr("xl/workbook.xml", WORKBOOK_XML)
        z.writestr("xl/_rels/workbook.xml.rels", RELS_XML)
        z.writestr("xl/styles.xml", STYLES_TEMPLATE.format(fills=fills))
        z.writestr("xl/worksheets/sheet1.xml", WORKSHEET_XML)
    return path


def run_audit(path: str, *flags: str) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, SCRIPT, path, *flags],
                          capture_output=True, text=True)


def violation_types(path: str) -> list:
    """Run the audit in --json mode and return the reported violation types."""
    result = run_audit(path, "--json")
    payload = json.loads(result.stdout)
    return [item["type"] for item in payload["violations"]]


class TestPatternFillDefaults(unittest.TestCase):
    def test_omitted_pattern_type_is_the_none_default(self):
        """The openpyxl form <patternFill/> must audit clean (issue #4990)."""
        with tempfile.TemporaryDirectory() as root:
            xlsx = build_xlsx(os.path.join(root, "default.xlsx"), DEFAULT_FILLS)
            result = run_audit(xlsx)
            self.assertEqual(result.returncode, 0, result.stdout)
            self.assertIn("Violations : 0", result.stdout)
            self.assertIn("PASS", result.stdout)
            self.assertNotIn("fills[0]", result.stdout)

    def test_empty_pattern_type_is_the_none_default(self):
        """An explicitly empty patternType="" is the same default."""
        with tempfile.TemporaryDirectory() as root:
            xlsx = build_xlsx(os.path.join(root, "empty-attr.xlsx"), EMPTY_ATTR_FILLS)
            result = run_audit(xlsx)
            self.assertEqual(result.returncode, 0, result.stdout)
            self.assertIn("PASS", result.stdout)

    def test_explicit_non_none_fill_still_fails(self):
        """A real fills[0] pattern keeps being reported."""
        with tempfile.TemporaryDirectory() as root:
            xlsx = build_xlsx(os.path.join(root, "solid.xlsx"), SOLID_FILLS)
            result = run_audit(xlsx)
            self.assertEqual(result.returncode, 1, result.stdout)
            self.assertIn("fills[0] patternType='solid', must be 'none'", result.stdout)
            self.assertIn("fills_0_corrupted", violation_types(xlsx))

    def test_default_fill_is_still_not_gray125(self):
        """The default is "none", so a default fills[1] must still fail."""
        with tempfile.TemporaryDirectory() as root:
            xlsx = build_xlsx(os.path.join(root, "no-gray125.xlsx"), NO_GRAY125_FILLS)
            result = run_audit(xlsx)
            self.assertEqual(result.returncode, 1, result.stdout)
            self.assertIn("fills[1] patternType='none', must be 'gray125'", result.stdout)
            self.assertIn("fills_1_corrupted", violation_types(xlsx))

    def test_gradient_fill_has_no_pattern_default(self):
        """A fill without <patternFill> is not silently read as "none"."""
        with tempfile.TemporaryDirectory() as root:
            xlsx = build_xlsx(os.path.join(root, "gradient.xlsx"), GRADIENT_FILLS)
            result = run_audit(xlsx)
            self.assertEqual(result.returncode, 1, result.stdout)
            self.assertIn("fills[0] patternType='', must be 'none'", result.stdout)

    @unittest.skipUnless(HAS_OPENPYXL, "openpyxl is not installed")
    def test_real_openpyxl_workbook_audits_clean(self):
        """End-to-end: a plain openpyxl workbook must not be rejected."""
        import openpyxl

        with tempfile.TemporaryDirectory() as root:
            xlsx = os.path.join(root, "openpyxl.xlsx")
            book = openpyxl.Workbook()
            sheet = book.active
            sheet["A1"] = "Revenue"
            sheet["B1"] = 1234.5
            book.save(xlsx)

            result = run_audit(xlsx)
            self.assertEqual(result.returncode, 0, result.stdout)
            self.assertIn("PASS", result.stdout)


if __name__ == "__main__":
    unittest.main()
